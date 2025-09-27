# word.py — exposes: extract_text(path) -> str
from typing import List
from docx import Document

def extract_text(path: str) -> str:
    """Return plain text from a DOCX (paragraphs + tables). Returns '' on failure."""
    try:
        doc = Document(path)
    except Exception:
        return ""

    parts: List[str] = []
    for p in doc.paragraphs:
        if p.text and p.text.strip():
            parts.append(p.text)

    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))

    return "\n".join(parts).strip()
